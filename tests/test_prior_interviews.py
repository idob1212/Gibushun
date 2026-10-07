"""Tests for the interview check: the admin syncs the unit's interview form
(the Google Forms "תגובות" Excel) and candidates whose name is in it are
flagged — on the admin home for every group, on a group's home page for its
own candidates, on the candidate page and in the Excel export.

Run with the pinned deps:
    python tests/test_prior_interviews.py
"""
import io
import os
from datetime import datetime

os.environ["DATABASE_URL"] = "sqlite:////tmp/test_prior_interviews.db"

import openpyxl  # noqa: E402
from sqlalchemy import event  # noqa: E402

from main import (app, db, User, Candidate, PriorInterview, InterviewSync,  # noqa: E402
                  DOCTOR_ID, name_key, parse_interview_file, prior_interview_matches)

app.config["WTF_CSRF_ENABLED"] = False

# The real form's columns, in its order, with its stray trailing spaces.
FORM_HEADER = ["חותמת זמן", "שם מראיין", "תעודת זהות מרואיין", "שם מלא מרואיין", "איזה גיבוש ביצע",
               "מדד תכנוני", "מהי המוטיבציה לשירות ביחידה ", "יכולת השתלבות בצוות", "חוסן מנטלי",
               "עד כמה מתאים לשרת ביחידה", 'האם החייל העלה בעיות ת"ש',
               "האם העלה בעיות רפואיות (במידה וכן נא לפרט", "האם אתה צופה שהחייל ידרג את היחידה ",
               "התרשמות כללית (נא להרחיב)", "דירוג המרואיין "]


def form_row(when, interviewer, name, gibush, rating, impression):
    return [when, interviewer, 123456789, name, gibush, 5, 5, 4, 4, 4, "לא", "בעיות בראייה", "כן",
            impression, rating]


ROWS = [
    form_row(datetime(2026, 5, 17, 9, 36, 16, 587000), "דנה מראיינת", "  יונתן  לוי ", "מטכל", "לקחת",
             "בחור רציני.\nשאל שאלות טובות"),
    form_row(datetime(2026, 5, 18, 11, 2, 0), "רון", "אביב מרדכי כהן", "שייטת", "בלית ברירה", "ביישן"),
    form_row(datetime(2026, 5, 19, 8, 0, 0), "רון", "נועם פרץ", "חובלים", "קו אדום", "לא מתאים"),
    form_row(datetime(2026, 5, 19, 9, 0, 0), "רון", 300000017, "x", "לקחת", "מספר במקום שם"),
    form_row(datetime(2026, 5, 20, 10, 0, 0), "טל", "ג'ורג' בן-חיים", "צוללות", "להתאבד", "מצוין"),
    form_row(datetime(2026, 5, 20, 12, 0, 0), "טל", "דניאל", "צוללות", "לקחת", "שם של מילה אחת"),
    [None] * len(FORM_HEADER),  # blank line at the end of the sheet
]


def build_xlsx(rows=ROWS, header=FORM_HEADER, summary_sheet=True):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "טופס מקורי"
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    if summary_sheet:
        # The real export has a second, narrower summary sheet; the form sheet wins.
        summary = workbook.create_sheet("סיכום עם ציונים")
        summary.append(["תעודת זהות מרואיין", "שם מלא מרואיין", "איזה גיבוש ביצע", "ציון", "דירוג המרואיין "])
        summary.append([111, "מישהו אחר לגמרי", "מטכל", 8, "לקחת"])
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


with app.app_context():
    db.drop_all()
    db.create_all()
    db.session.add_all([
        User(id=0, name="admin", password="x", mitam_num=0, sprint_num=0, crawl_num=0, alonka_num=0),
        User(id=1, name="גבעתי", password="a", mitam_num=1, sprint_num=1, crawl_num=1, alonka_num=1),
        User(id=2, name="גולני", password="b", mitam_num=1, sprint_num=1, crawl_num=1, alonka_num=1),
        User(id=DOCTOR_ID, name="רופא", password="d", mitam_num=0, sprint_num=0, crawl_num=0, alonka_num=0),
    ])
    db.session.add_all([
        Candidate(id="1/1", group_id=1, name="לוי יונתן"),        # same words, other order
        Candidate(id="1/2", group_id=1, name="אביב כהן"),          # middle name left out
        Candidate(id="1/3", group_id=1, name="מגובש 3"),           # placeholder
        Candidate(id="1/4", group_id=1, name="דניאל"),             # one word
        Candidate(id="1/5", group_id=1, name="נועם פרץ", status="פרש"),  # retired
        Candidate(id="2/1", group_id=2, name="ג׳ורג׳ בן חיים"),    # other group, other quotes/dash
        Candidate(id="2/2", group_id=2, name="איתי שמש"),          # not interviewed
    ])
    db.session.commit()

client = app.test_client()


def login_as(user_id):
    with client.session_transaction() as sess:
        sess["_user_id"] = str(user_id)
        sess["_fresh"] = True


def html_of(resp):
    return resp.get_data(as_text=True)


def upload(data, filename="ראיונות (תגובות).xlsx", next_url="/1/"):
    return client.post("/interviews/sync", data={"file": (io.BytesIO(data), filename), "next": next_url},
                       content_type="multipart/form-data")


# 1) Names compare without niqqud, quotes, dashes, final letters or spacing
assert name_key("  ג׳ורג׳ בן-חיים ") == name_key("ג'ורג' בן חיים") == "גורג בנ חיימ"
assert name_key("אַבְרָהָם כֹּהֵן") == name_key("אברהם כהן")
assert name_key('יצחק "איציק" לוי') == "יצחק איציק לוי"
assert name_key("Noam  COHEN") == "noam cohen"
print("OK: names are normalised before comparing")

# 2) Parsing the form's export
rows, skipped = parse_interview_file(build_xlsx())
assert len(rows) == 5 and skipped == 1, (len(rows), skipped)   # the ID-number row is skipped
first = rows[0]
assert first["name"] == "יונתן לוי" and first["interviewer"] == "דנה מראיינת"
assert first["interviewed_at"] == datetime(2026, 5, 17, 9, 36, 16)
assert first["rating"] == "לקחת" and first["gibush"] == "מטכל"
assert first["impression"] == "בחור רציני.\nשאל שאלות טובות"
assert "id" not in first and not any("123456789" in str(v) for v in first.values())  # no ID numbers kept
# The same answers as CSV (Google Sheets "download as CSV"), UTF-8 with BOM and Windows-1255.
csv_text = ",".join(f'"{h}"' for h in FORM_HEADER) + "\n" + \
    '"17/05/2026 9:36:16","דנה","1","יונתן לוי","מטכל","5","5","4","4","4","לא","לא","כן","טוב","לקחת"\n'
for encoding in ("utf-8-sig", "cp1255"):
    csv_rows, _ = parse_interview_file(csv_text.encode(encoding))
    assert len(csv_rows) == 1 and csv_rows[0]["name_key"] == "יונתנ לוי", csv_rows
    assert csv_rows[0]["interviewed_at"] == datetime(2026, 5, 17, 9, 36, 16)
for bad, message in [(b"\xd0\xcf\x11\xe0" + b"\0" * 100, "xls"),
                     (build_xlsx(rows=[], header=["שם", "טלפון"], summary_sheet=False), "שם מלא מרואיין")]:
    try:
        parse_interview_file(bad)
        raise AssertionError("bad file accepted")
    except ValueError as error:
        assert message in str(error), str(error)
print("OK: the form export parses (xlsx and csv); bad files are explained")

# 3) Only the admin syncs
for user_id in (1, DOCTOR_ID):
    login_as(user_id)
    assert upload(build_xlsx()).status_code == 403
with app.app_context():
    assert PriorInterview.query.count() == 0
print("OK: groups and staff cannot sync")

# 4) Admin sync: summary, popup on the next page, flags on the rows
login_as(0)
html = html_of(client.get("/1/"))
assert "בדיקת ראיונות קודמים" in html and "העלאת קובץ ראיונות" in html and 'id="pi-sheet"' not in html
resp = upload(build_xlsx())
assert resp.status_code == 302 and resp.headers["Location"].endswith("/1/")
html = html_of(client.get("/1/"))
assert "סונכרנו 5 ראיונות" in html and "שורה אחת בלי שם דולגה" in html and "3 מגובשים כבר רואיינו" in html
assert 'id="pi-sheet"' in html and "data-open-now" in html           # pops up right after the sync
assert "קבוצה 1" in html and "קבוצה 2" in html                       # admin sees every group
assert "שם דומה: אביב מרדכי כהן" in html                             # middle name: a possible match
assert "לקחת" in html and "להתאבד" in html                           # admin sees the rating
with app.app_context():
    matches = prior_interview_matches(Candidate.query.filter(Candidate.group_id > 0).all())
    assert set(matches) == {"1/1", "1/2", "1/5", "2/1"}, set(matches)  # 1/5 is retired, see below
    assert matches["1/1"]["exact"] and matches["2/1"]["exact"] and not matches["1/2"]["exact"]
    assert InterviewSync.query.count() == 1 and PriorInterview.query.count() == 5
html = html_of(client.get("/1/"))
assert "data-open-now" not in html and 'id="pi-sheet"' in html       # later visits: only if unseen
assert "סנכרון קובץ מעודכן" in html and "5 ראיונות בטופס" in html
print("OK: admin sync reports, pops up once, lists matches from every group")

# 5) Retired candidates, placeholders and one-word names are not flagged
assert html.count('class="pi-note') == 2, html.count('class="pi-note')  # 1/1 and 1/2 in group 1's list
keys = [m["key"] for m in matches.values()]
assert len(set(keys)) == len(keys)
print("OK: placeholders, single words and retired candidates are left out")

# 6) Re-syncing the same (or a newer) export updates instead of duplicating
newer = list(ROWS)
newer[0] = form_row(datetime(2026, 5, 17, 9, 36, 16), "דנה מראיינת", "יונתן לוי", "מטכל", "בלית ברירה", "עודכן")
newer.append(form_row(datetime(2026, 5, 21, 9, 0), "רון", "איתי שמש", "שייטת", "לקחת", "חדש"))
html = html_of(client.get(upload(build_xlsx(newer)).headers["Location"]))
assert "סונכרנו 6 ראיונות (אחד חדש)" in html, html[html.find("סונכרנו"):][:120]
with app.app_context():
    assert PriorInterview.query.count() == 6
    updated = PriorInterview.query.filter_by(name_key=name_key("יונתן לוי")).one()
    assert updated.rating == "בלית ברירה" and updated.impression == "עודכן"
print("OK: re-sync is idempotent and picks up edited answers")

# 7) A group sees only its own candidates, without the rating or impression
login_as(1)
html = html_of(client.get("/"))
assert 'id="pi-sheet"' in html and "2 מגובשים כבר רואיינו" in html
assert "data-open-now" not in html
assert "לוי יונתן" in html and "ג׳ורג׳" not in html and "איתי שמש" not in html
assert "בלית ברירה" not in html and "עודכן" not in html
assert "דנה מראיינת" in html and "17/05/2026" in html
html = html_of(client.get("/candidate/1/1"))
assert "רואיין ביחידה" in html and "דנה מראיינת" in html and "מטכל" in html
assert "בלית ברירה" not in html and "עודכן" not in html
html = html_of(client.get("/candidate/1/2"))
assert "ייתכן שרואיין ביחידה" in html and "אביב מרדכי כהן" in html
login_as(2)
html = html_of(client.get("/"))
assert "2 מגובשים כבר רואיינו" in html and "ג׳ורג׳ בן חיים" in html and "לוי יונתן" not in html
print("OK: each group sees its own matches; rating stays with the admin")

# 8) The admin's candidate page shows the full answer
login_as(0)
html = html_of(client.get("/candidate/1/1"))
assert "בלית ברירה" in html and "עודכן" in html and "התרשמות" in html
print("OK: admin candidate page shows rating and impression")

# 9) Naming a placeholder after the sync flags it too
with app.app_context():
    Candidate.query.get("1/3").name = "נועם פרץ"
    db.session.commit()
login_as(1)
html = html_of(client.get("/"))
assert "3 מגובשים כבר רואיינו" in html
print("OK: a candidate named after the sync is flagged")

# 10) The admin's Excel export carries the match
login_as(0)
resp = client.get("/download-sheet/")
assert resp.status_code == 200
sheet = openpyxl.load_workbook(io.BytesIO(resp.data), read_only=True).worksheets[0]
table = [list(r) for r in sheet.iter_rows(values_only=True)]
column = table[0].index("ראיון קודם ביחידה")
by_id = {r[0]: r[column] for r in table[1:]}
assert "17/05/2026" in by_id["1/1"] and "דנה מראיינת" in by_id["1/1"] and "בלית ברירה" in by_id["1/1"]
assert "שם דומה" in by_id["1/2"] and "רון" in by_id["2/2"] and not by_id["1/4"]
assert table[0][-1] == "ציון סופי משוקלל ראיון וגיבוש"  # still the last column
print("OK: export has a 'prior interview' column")

# 11) Home page stays cheap with a synced list
statements = []


def count(conn, cursor, statement, *args):
    statements.append(statement)


login_as(1)
client.get("/")
with app.app_context():
    event.listen(db.engine, "before_cursor_execute", count)
    client.get("/")
    event.remove(db.engine, "before_cursor_execute", count)
assert len(statements) <= 14, len(statements)
print(f"OK: group home with the check = {len(statements)} queries")

# 12) Bad uploads are refused with a message; nothing changes
login_as(0)
for data, name, message in [(b"hello", "a.txt", "שם מלא מרואיין"),
                            (b"\xd0\xcf\x11\xe0" + b"\0" * 64, "old.xls", "xls"),
                            (b"PK\x03\x04garbage", "broken.xlsx", "הקובץ פגום")]:
    html = html_of(client.get(upload(data, name).headers["Location"]))
    assert message in html, (name, message)
html = html_of(client.get(client.post("/interviews/sync", data={"next": "/1/"}).headers["Location"]))
assert "בחרו את קובץ" in html
big = b"PK\x03\x04" + b"0" * (5 * 1024 * 1024 + 10)
assert "גדול מדי" in html_of(client.get(upload(big).headers["Location"]))
with app.app_context():
    assert PriorInterview.query.count() == 6
print("OK: wrong, broken and oversized files are refused")

# 13) Clearing the list removes every flag
assert client.get("/interviews/clear").status_code == 405  # POST only
resp = client.post("/interviews/clear", data={"next": "/1/"})
html = html_of(client.get(resp.headers["Location"]))
assert "רשימת הראיונות נמחקה" in html and 'id="pi-sheet"' not in html and "העלאת קובץ ראיונות" in html
with app.app_context():
    assert PriorInterview.query.count() == 0 and InterviewSync.query.count() == 0
login_as(1)
assert 'id="pi-sheet"' not in html_of(client.get("/"))
print("OK: clearing the list removes the alerts")

print("ALL PRIOR INTERVIEW TESTS PASSED")
